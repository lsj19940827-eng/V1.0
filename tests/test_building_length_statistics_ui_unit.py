"""长度明细、类型汇总、复制排版及工程恢复使用同一数据源。"""

import copy
import math
import os

import pytest

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
os.environ.setdefault('QTWEBENGINE_DISABLE_SANDBOX', '1')

from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QApplication, QPushButton
from app_渠系计算前端.water_profile.water_profile_dialogs import BuildingLengthDialog
from app_渠系计算前端.water_profile.panel import WaterProfilePanel
from 推求水面线.core.calculator import WaterProfileCalculator
from 推求水面线.core.length_statistics import summarize_length_records
from 推求水面线.models.data_models import ProjectSettings
from test_building_length_statistics_unit import _connection_case, _node


@pytest.fixture(scope='module')
def app():
    return QApplication.instance() or QApplication([])


def _dispose(app, *widgets):
    app.processEvents()
    for widget in widgets:
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    app.processEvents()


def test_detail_summary_and_formatted_copy_include_the_same_transition_lengths(app, monkeypatch):
    import app_渠系计算前端.water_profile.water_profile_dialogs as dialog_module
    monkeypatch.setattr(dialog_module, 'fluent_info', lambda *args, **kwargs: None)
    records = WaterProfileCalculator(ProjectSettings()).calculate_building_lengths(_connection_case())
    stale_summary = [{'structure_type': '渐变段', 'count': 1, 'total_length': 999}]
    detail = BuildingLengthDialog(None, records, 230, stale_summary, station_prefix='茶支')
    formatted = detail.formatted_view
    try:
        assert detail._type_summary == summarize_length_records(records)
        detail_total = math.fsum(float(detail.detail_table.item(r, 3).text()) for r in range(detail.detail_table.rowCount()))
        summary_total = float(detail.summary_table.item(detail.summary_table.rowCount() - 1, 3).text())
        assert detail_total == summary_total == 230
        rows = formatted._table_data
        assert math.fsum(float(row[3]) for row in rows if row[3]) == 230
        assert next(float(row[5]) for row in rows if row[4] == '总长度') == 230
        assert sum('渐变段' in row[0] for row in rows) == 2
        assert '渐变段' in formatted._generate_tsv_text()
        assert detail.tab_widget.currentIndex() == 0
        assert detail.tab_widget.tabText(0) == '排版表预览(Excel)'
        assert detail.detail_table.item(0, 1).text() == '茶支0+000.000'
        buttons = {button.text(): button for button in detail.findChildren(QPushButton)}
        assert set(buttons) == {'复制到剪贴板', '关闭'}
        # 选中局部仍复制完整排版表，页签切换后复制内容跟随当前表格。
        detail.detail_table.setCurrentCell(0, 0)
        buttons['复制到剪贴板'].click()
        assert QApplication.clipboard().text() == formatted._generate_tsv_text()
        assert all(len(line.split('\t')) == 6 for line in QApplication.clipboard().text().splitlines())
        detail.tab_widget.setCurrentIndex(1)
        buttons['复制到剪贴板'].click()
        copied = [line.split('\t') for line in QApplication.clipboard().text().splitlines()]
        assert copied[0] == ['序号', '结构类型', '数量', '累计长度(m)']
        assert float(copied[-1][3]) == 230
        assert len(copied) == detail.summary_table.rowCount() + 1
    finally:
        _dispose(app, detail)


def test_displayed_lengths_close_to_displayed_stations_and_summary(app):
    boundaries = [0.0004, 0.0008, 1.0002, 1000.0005]
    records = [{'name': f'明渠{i}', 'structure_type': '明渠-矩形', 'length': b - a,
                'start_station': a, 'end_station': b, 'node_count': 2}
               for i, (a, b) in enumerate(zip(boundaries, boundaries[1:]))]
    dialog = BuildingLengthDialog(None, records, boundaries[-1] - boundaries[0])
    formatted = dialog.formatted_view
    try:
        lengths = []
        for row in range(dialog.detail_table.rowCount()):
            if not dialog.detail_table.item(row, 3).text():
                continue
            length = float(dialog.detail_table.item(row, 3).text())
            stations = [dialog.detail_table.item(row, col).text().split('+') for col in (1, 2)]
            start, end = [float(km) * 1000 + float(metres) for km, metres in stations]
            assert length == pytest.approx(end - start, abs=1e-9)
            assert f"{records[row]['length']:.12g}" in dialog.detail_table.item(row, 3).toolTip()
            lengths.append(length)
        assert sum(lengths) == pytest.approx(float(dialog.summary_table.item(1, 3).text()))
        assert formatted._format_station(999.9996) == '1+000.000'
    finally:
        _dispose(app, dialog)


def test_project_reopen_rebuilds_lengths_and_replaces_previous_project_cache(app):
    panel, restored = WaterProfilePanel(), WaterProfilePanel()
    try:
        panel._settings = ProjectSettings(channel_name='统计往返', start_station=0)
        panel.nodes = [_node('甲', '隧洞-圆形', 0, '进'), _node('甲', '隧洞-圆形', 123.456789, '出')]
        panel._update_table_from_nodes_full(panel.nodes, '统计支')
        panel._rebuild_calculation_summary_state(panel.nodes)
        expected = copy.deepcopy(panel._last_building_lengths)
        payload = panel.to_project_dict()
        restored._last_channel_total_length = 999
        restored._last_building_lengths = [{'name': '旧工程', 'length': 999}]
        restored.from_project_dict(payload, skip_dirty_signal=True)
        assert restored._last_channel_total_length == 123.456789
        assert restored._last_building_lengths == expected
        assert restored._build_nodes_from_table()[-1].station_MC == 123.456789
        assert restored.btn_building_stats.isEnabled()
        # 表格快照可以比节点模型更新，恢复统计时必须以当前表格为准。
        edited_payload = copy.deepcopy(payload)
        edited_payload['node_table_rows'][-1][15] = '统计支0+130.125'
        edited_payload['node_numeric_display_values'][-1].pop('15', None)
        restored.from_project_dict(edited_payload, skip_dirty_signal=True)
        assert restored._last_channel_total_length == 130.125
        restored.from_project_dict({}, skip_dirty_signal=True)
        assert restored._last_building_lengths == []
        assert restored._last_type_summary == []
        assert not restored.btn_building_stats.isEnabled()
    finally:
        _dispose(app, panel, restored)


def test_invalid_new_nodes_clear_previous_statistics(app):
    panel = WaterProfilePanel()
    try:
        panel._settings = ProjectSettings()
        panel._refresh_building_length_state(_connection_case())
        assert panel._last_building_lengths
        panel._rebuild_calculation_summary_state([_node('甲', '隧洞-圆形', 100), _node('甲', '隧洞-圆形', 90)])
        assert panel._last_building_lengths == []
        assert panel._last_type_summary == []
        assert '桩号逆序' in panel._last_building_stats_error
        assert not panel.btn_building_stats.isEnabled()
    finally:
        _dispose(app, panel)


def test_repeated_transition_edits_use_current_gap_and_refresh_saved_connection_length(app):
    panel = WaterProfilePanel()
    try:
        nodes = _connection_case()
        nodes[2].transition_length = 8
        nodes[2].transition_length_override_m = 8
        panel._update_table_from_nodes_full(nodes)
        panel._refresh_building_length_state(nodes)
        assert nodes[3].stat_length == 17
        assert panel._build_nodes_from_table()[3].stat_length == 17
        for row, expected in ((2, 25), (4, 22)):
            context = {'node': nodes[row], 'nodes': nodes, 'row_idx': row}
            assert panel._get_transition_length_override_upper_bound(context) == expected
    finally:
        _dispose(app, panel)


def test_word_length_table_matches_current_detail_and_displayed_total(app, tmp_path):
    from docx import Document
    panel = WaterProfilePanel()
    try:
        panel._settings = ProjectSettings(channel_name='统计导出验证')
        panel.calculated_nodes = _connection_case()
        panel._last_building_lengths = [{'name': '过期缓存', 'length': 999}]
        target = tmp_path / 'length-report.docx'
        panel._build_word_report(str(target))
        document = Document(target)
        table = next(t for t in document.tables if len(t.columns) == 6 and t.cell(0, 0).text == '序号')
        rows = [[cell.text for cell in row.cells] for row in table.rows[1:]]
        assert math.fsum(float(row[3]) for row in rows[:-1]) == float(rows[-1][3]) == 230
        assert sum(row[2] == '渐变段' for row in rows[:-1]) == 2
        assert all('过期缓存' not in row for row in rows)
    finally:
        _dispose(app, panel)


def test_loss_edit_details_undo_and_redo_preserve_values_below_display_precision(app, monkeypatch):
    import app_渠系计算前端.water_profile.panel as panel_module
    monkeypatch.setattr(panel_module.InfoBar, 'success', lambda *args, **kwargs: None)
    panel = WaterProfilePanel()
    try:
        nodes = [_node('', '明渠-矩形', 0), _node('', '明渠-矩形', 100)]
        upstream_level, friction, reserve = 100.123456, 0.1234567, 0.0001234
        nodes[0].water_level = upstream_level
        nodes[1].water_level = upstream_level - friction
        nodes[1].head_loss_friction = nodes[1].head_loss_total = nodes[1].head_loss_cumulative = friction
        panel.calculated_nodes = nodes
        panel.start_wl_edit.setText(str(upstream_level))
        panel._update_table_from_nodes_full(nodes)
        snapshot = panel._snapshot_editable_cols()
        panel._updating_cells = True
        panel.node_table.item(1, 36).setText(str(reserve))
        panel._recalc_downstream(1)
        panel._sync_losses_from_table()
        expected = upstream_level - friction - reserve
        assert panel.calculated_nodes[1].water_level == pytest.approx(expected, abs=1e-12)
        assert panel._build_nodes_from_table()[1].water_level == pytest.approx(expected, abs=1e-12)
        panel._node_table_undo_stack.clear()
        panel._node_table_redo_stack.clear()
        panel._loss_undo_stack = [snapshot]
        panel._loss_redo_stack.clear()
        panel._undo_loss_edit()
        assert panel._build_nodes_from_table()[1].water_level == pytest.approx(upstream_level - friction, abs=1e-12)
        panel._redo_loss_edit()
        assert panel.calculated_nodes[1].water_level == pytest.approx(expected, abs=1e-12)
        assert panel._build_nodes_from_table()[1].water_level == pytest.approx(expected, abs=1e-12)
    finally:
        _dispose(app, panel)


def test_transition_detail_sync_keeps_length_beyond_three_display_decimals(app):
    panel = WaterProfilePanel()
    try:
        nodes = _connection_case()
        nodes[2].transition_length = 7.705700872036023
        panel._update_table_from_nodes_full(nodes)
        assert panel.node_table.item(2, 32).text() == '7.706'
        assert panel._get_transition_length_cell_value(2) == (True, 7.705700872036023)
        panel._sync_transition_lengths_from_table(nodes)
        assert nodes[2].transition_length == 7.705700872036023
    finally:
        _dispose(app, panel)
