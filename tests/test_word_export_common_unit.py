# -*- coding: utf-8 -*-
"""全程序 Word 导出的版本、分页和普通表格回归。"""

import ast
import importlib.util
import os
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import version


def _load_export():
    """独立加载真实导出实现，避免其他面板测试的全局替身干扰。"""
    path = ROOT / "app_渠系计算前端/export_utils.py"
    spec = importlib.util.spec_from_file_location("word_export_common_under_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _meta():
    return SimpleNamespace(
        project_name="导出格式核验工程", design_stage="施工图", product_level="三级",
        record_number="", specialty="水工", calculator="", checker="", reviewer="",
        approver="", volume_current="1", volume_total="1", basic_info="",
        mandatory_clause="无", extra_references=[],
    )


def _report(export, **kwargs):
    return export.create_engineering_report_doc(
        _meta(), "有压管道水力计算书", "有压管道水力计算（球墨铸铁管）",
        "按设计流量和管材校核水力条件。", ["《灌溉与排水工程设计标准》(GB 50288-2018)"],
        **kwargs,
    )


def _assert_plain_table(table):
    """检查存储的无底纹语义，而不是把白色填充视作透明。"""
    assert table.style.name == "Table Grid"
    assert not table._tbl.xpath('.//w:highlight')
    for shade in table._tbl.xpath('.//w:shd'):
        assert shade.get(qn('w:val')) == 'nil'
        assert shade.get(qn('w:fill')) is None
        assert shade.get(qn('w:themeFill')) is None
    for color in table._tbl.xpath('.//w:rPr/w:color'):
        assert dict(color.attrib) == {qn('w:val'): '000000'}
    borders = table._tbl.tblPr.find(qn('w:tblBorders'))
    assert len(borders) == 6
    for border in borders:
        assert border.get(qn('w:val')) == 'single'
        assert border.get(qn('w:color')) == '000000'


@pytest.mark.parametrize("future_version", ["1.4.4", "2.7.3.1"])
@pytest.mark.parametrize("kind", ["template", "fallback", "batch"])
def test_report_version_follows_single_release_source(monkeypatch, tmp_path, kind, future_version):
    monkeypatch.setattr(version, 'APP_VERSION', future_version)
    export = _load_export()
    if kind == 'batch':
        doc = export.create_styled_doc("批量水力计算书")
    else:
        if kind == 'fallback':
            monkeypatch.setattr(export, '_TEMPLATE_PATH', str(tmp_path / 'absent.docx'))
        doc = _report(export)
    target = tmp_path / f'{kind}.docx'
    doc.save(target)
    reopened = Document(target)
    text = '\n'.join(reopened._element.xpath('.//w:t/text()'))
    assert f'V{future_version}' in text
    assert 'V1.0' not in text
    for table in reopened.tables:
        _assert_plain_table(table)


@pytest.mark.parametrize("relative_path", [
    'open_channel/panel.py', 'aqueduct/panel.py', 'tunnel/panel.py',
    'culvert/panel.py', 'pressure_pipe/panel.py', 'siphon/panel.py',
    'water_profile/panel.py', 'spillway_steep_chute/report_export.py',
])
def test_module_program_text_uses_release_version_and_retains_description(relative_path):
    """核对各实际导出入口的文字，防止任一面板继续覆盖公共默认版本。"""
    source = ROOT / 'app_渠系计算前端' / relative_path
    tree = ast.parse(source.read_text(encoding='utf-8'))
    assert any(isinstance(node, ast.ImportFrom) and node.module == 'version'
               and any(name.name == 'APP_VERSION' for name in node.names) for node in tree.body)
    calls = [node for node in ast.walk(tree) if isinstance(node, ast.Call)
             and isinstance(node.func, ast.Name) and node.func.id == 'create_engineering_report_doc']
    assert calls
    for call in calls:
        expr = next(item.value for item in call.keywords if item.arg == 'calc_program_text')
        scope = dict(APP_VERSION='9.8.7', desc='有压管道计算', content_desc='断面水力计算',
                     stype='圆形', method='试算')
        text = eval(compile(ast.Expression(expr), str(source), 'eval'), {'__builtins__': {}}, scope)
        assert text.splitlines()[0] == '渠系建筑物水力计算系统 V9.8.7'
        assert len(text.splitlines()) == 2
        assert text.splitlines()[1]


@pytest.mark.parametrize('kind', ['data', 'result', 'params'])
def test_common_tables_are_plain_and_keep_values_after_save(tmp_path, kind):
    export = _load_export()
    doc = Document()
    rows = [('推荐规格', 'DN800'), ('计算内径', '810.8 mm'), ('状态', '★选中')]
    if kind == 'data':
        export.doc_add_styled_table(doc, ['项目', '值'], rows,
                                    highlight_col=1, highlight_val='★选中', with_full_border=False)
    elif kind == 'result':
        export.doc_add_result_table(doc, rows)
    else:
        export.doc_add_param_table(doc, rows)
    doc.save(tmp_path / 'table.docx')
    table = Document(tmp_path / 'table.docx').tables[0]
    _assert_plain_table(table)
    assert [[c.text for c in row.cells] for row in table.rows[1:]] == [list(row) for row in rows]
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')


def test_plain_style_clears_existing_fill_and_preserves_merges_and_formula():
    export = _load_export()
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Light Shading Accent 1'
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = '合并表头'
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:fill="FFFFFF" w:themeFill="accent1"/>'
    ))
    cell.paragraphs[0].runs[0]._r.get_or_add_rPr().append(parse_xml(
        f'<w:color {nsdecls("w")} w:val="C9A96E" w:themeColor="accent1"/>'
    ))
    formula = parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>A=1</m:t></m:r></m:oMath>'
    )
    table.cell(1, 0).paragraphs[0]._p.append(formula)
    formula_before = etree.tostring(formula)
    export.doc_apply_plain_table_style(table)
    _assert_plain_table(table)
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(0, 0).text == '合并表头'
    assert etree.tostring(formula) == formula_before
    xml_before = table._tbl.xml
    export.doc_apply_plain_table_style(table)
    assert table._tbl.xml == xml_before


def test_template_section_break_preserves_numbering_without_empty_cover_tail():
    export = _load_export()
    original = Document(export._TEMPLATE_PATH)
    doc = _report(export)
    original_sections = original._element.xpath('.//w:sectPr')
    sections = doc._element.xpath('.//w:sectPr')
    assert [section.xml for section in sections] == [section.xml for section in original_sections]
    toc = doc._element.body.find(qn('w:sdt'))
    cover_last = toc.getprevious()
    assert cover_last.tag == qn('w:p')
    assert '月' in ''.join(cover_last.xpath('.//w:t/text()'))
    assert cover_last.xpath('./w:pPr/w:sectPr')
    assert not cover_last.xpath('.//w:br[@w:type="page"]')
    title = toc.find(qn('w:sdtContent')).find(qn('w:p'))
    assert title.xpath('./w:pPr/w:outlineLvl/@w:val') == ['9']
    assert len(doc.tables) == len(original.tables)
    for before, after in zip(original.tables, doc.tables):
        assert before._tbl.xpath('.//w:gridSpan/@w:val') == after._tbl.xpath('.//w:gridSpan/@w:val')
        assert before._tbl.xpath('.//w:vMerge/@w:val') == after._tbl.xpath('.//w:vMerge/@w:val')


@pytest.mark.skipif(os.environ.get('RUN_WORD_COM_TESTS') != '1', reason='按需启用 Microsoft Word 原生分页验证')
def test_word_native_cover_and_toc_are_adjacent(tmp_path):
    import win32com.client

    docx = _report(_load_export())
    target = tmp_path / 'native-layout.docx'
    docx.save(target)
    app = win32com.client.DispatchEx('Word.Application')
    app.Visible = False
    app.DisplayAlerts = 0
    opened = None
    try:
        opened = app.Documents.Open(str(target), ReadOnly=True, AddToRecentFiles=False)
        opened.Fields.Update()
        for index in range(1, opened.TablesOfContents.Count + 1):
            opened.TablesOfContents(index).Update()
        opened.Repaginate()
        count = opened.ComputeStatistics(2)
        pages = []
        for page in range(1, count + 1):
            start = opened.GoTo(1, 1, page).Start
            end = opened.GoTo(1, 1, page + 1).Start if page < count else opened.Content.End
            text = opened.Range(start, end).Text.replace('\x07', '').strip()
            assert text, f'第 {page} 页不应为空白页'
            pages.append(text)
        cover_index = next(i for i, text in enumerate(pages) if '四川水发勘测设计研究有限公司' in text)
        toc_index = next(i for i, text in enumerate(pages) if text.startswith('目录'))
        assert toc_index == cover_index + 1
        assert pages[toc_index].count('目录') == 1
        assert opened.GoTo(1, 1, toc_index + 2).Information(1) == 1
    finally:
        if opened is not None:
            opened.Close(False)
            opened = None
        app.Quit()
