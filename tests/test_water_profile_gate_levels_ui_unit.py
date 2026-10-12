"""真实表3的闸双行水位、保存、人工改损失及撤销回归。"""

import copy
import json
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("CODEX_FORCE_QTEXTBROWSER", "1")

import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication

from app_渠系计算前端.water_profile import panel as panel_module
from app_渠系计算前端.batch.panel import BatchPanel
from 推求水面线.core.calculator import WaterProfileCalculator
from 推求水面线.models.data_models import ChannelNode, ProjectSettings
from 推求水面线.models.enums import StructureType


@pytest.fixture
def panel(monkeypatch):
    app = QApplication.instance() or QApplication([])
    monkeypatch.setattr(BatchPanel, "_load_user_prefs", lambda self: None)
    for name in ("warning", "error", "info", "success"):
        monkeypatch.setattr(panel_module.InfoBar, name, lambda *args, **kwargs: None)
    value = panel_module.WaterProfilePanel()
    value.start_wl_edit.setText("100")
    value.design_flow_edit.setText("2")
    value._section_sync_ready = True
    value._transition_topology_prepared = True
    nodes = []
    for index, structure in enumerate((StructureType.MINGQU_RECTANGULAR, StructureType.CHECK_GATE, StructureType.MINGQU_RECTANGULAR)):
        node = ChannelNode(name=("上游明渠", "中间节制闸", "下游明渠")[index], structure_type=structure, flow_section="1", x=index * 100, flow=2)
        if index != 1:
            node.water_depth = 1.2
            node.structure_height = 2.5
            node.slope_i = 0.001
            node.section_params = {"B": 2, "h": 1.2, "H": 2.5}
        nodes.append(node)
    settings = ProjectSettings(start_water_level=100, design_flows=[2])
    # 100 m 来水区间损失 0.1 m，过闸另计 0.1 m；无台阶渠底按底坡连续。
    nodes = WaterProfileCalculator(settings).calculate_all(nodes)
    value.nodes = nodes
    value.calculated_nodes = nodes
    value._settings = settings
    value._update_table_from_nodes_full(nodes)
    yield value
    value.deleteLater()
    app.processEvents()


def test_gate_uses_original_water_level_cell_and_keeps_44_columns(panel):
    assert panel.node_table.columnCount() == 44
    assert len(panel_module.NODE_ALL_HEADERS) == 44
    assert panel.node_table.item(1, 41).text() == "闸前 99.900\n闸后 99.800"
    assert panel.node_table.item(1, 42).text() == "98.700"
    assert panel.node_table.item(1, 43).text() == "101.200"
    assert not (panel.node_table.item(1, 41).flags() & Qt.ItemIsEditable)
    assert panel._read_table_numeric_value(1, 41) == pytest.approx(99.8)
    assert panel.node_table.rowHeight(1) >= 2 * panel.node_table.fontMetrics().height() + 10
    assert panel.node_table._frozen_view.rowHeight(1) == panel.node_table.rowHeight(1)
    assert panel._build_nodes_from_table()[1].water_level == pytest.approx(99.8)
    assert panel._read_table_numeric_value(1, 35) == pytest.approx(0.1)


def test_gate_loss_edit_undo_redo_preserves_both_sides_and_downstream(panel):
    panel.node_table.setCurrentCell(1, 37)
    panel.node_table.item(1, 37).setText("0.0250")
    assert panel.node_table.item(1, 41).text() == "闸前 99.900\n闸后 99.875"
    assert panel.calculated_nodes[1].gate_upstream_water_level == pytest.approx(99.9)
    assert panel.calculated_nodes[1].gate_downstream_water_level == pytest.approx(99.875)
    assert panel._read_table_numeric_value(2, 41) == pytest.approx(99.775)
    assert panel.node_table.item(1, 42).text() == "98.700"
    assert panel.node_table.item(1, 43).text() == "101.200"
    panel._undo_loss_edit()
    assert panel.node_table.item(1, 41).text() == "闸前 99.900\n闸后 99.800"
    assert panel.calculated_nodes[1].gate_downstream_water_level == pytest.approx(99.8)
    assert panel._read_table_numeric_value(2, 41) == pytest.approx(99.7)
    assert panel.node_table.item(1, 42).text() == "98.700"
    assert panel.node_table.item(1, 43).text() == "101.200"
    panel._redo_loss_edit()
    assert panel.node_table.item(1, 41).text() == "闸前 99.900\n闸后 99.875"
    assert panel.calculated_nodes[1].gate_downstream_water_level == pytest.approx(99.875)
    assert panel._read_table_numeric_value(2, 41) == pytest.approx(99.775)
    assert panel.node_table.item(1, 42).text() == "98.700"
    assert panel.node_table.item(1, 43).text() == "101.200"


def test_gate_project_roundtrip_preserves_precise_sides_without_extra_columns(panel):
    gate = panel.calculated_nodes[1]
    gate.water_level = gate.gate_downstream_water_level = 99.912345678901
    gate.gate_upstream_water_level = 100.012345678901
    panel._update_table_from_nodes_full(panel.calculated_nodes)
    saved = json.loads(json.dumps(panel.to_project_dict()))
    assert saved["table_state"]["column_count"] == 44
    panel.from_project_dict(saved, skip_dirty_signal=True)
    decoded = panel._build_nodes_from_table()[1]
    assert decoded.gate_upstream_water_level == 100.012345678901
    assert decoded.gate_downstream_water_level == 99.912345678901
    assert decoded.water_level == 99.912345678901
    assert panel.node_table.item(1, 41).text() == "闸前 100.012\n闸后 99.912"
    # 显式水位精度完整保存；无台阶渠底依几何与底坡确定，不随过闸损失移动。
    assert decoded.bottom_elevation == pytest.approx(98.7)
    assert decoded.top_elevation == pytest.approx(101.2)
    assert decoded.terminal_gate_backfill_details["donor_name"] == "上游明渠"


def test_legacy_project_recovers_gate_display_without_recalculating_losses(panel):
    saved = copy.deepcopy(panel.to_project_dict())
    for row in saved["table_state"]["rows"]:
        data = row["cell_data"].get("0", {})
        data.pop("_gate_upstream_water_level", None)
        data.pop("_gate_downstream_water_level", None)
        data.pop("_gate_elevation_reference_details", None)
    gate_row = saved["table_state"]["rows"][1]
    gate_row["cells"][41] = "99.900"
    gate_row["cell_data"]["0"]["_numeric_display_values"]["41"] = {"text": "99.900", "value": 99.9}
    for column in (42, 43):
        gate_row["cells"][column] = "-"
        gate_row["cell_data"]["0"]["_numeric_display_values"][str(column)] = {"text": "-", "value": 0}
    for group_name in ("nodes", "calculated_nodes"):
        for node in saved.get(group_name, []):
            node.pop("gate_upstream_water_level", None)
            node.pop("gate_downstream_water_level", None)
            node.pop("terminal_gate_backfill_details", None)
        saved[group_name][1]["bottom_elevation"] = 0
        saved[group_name][1]["top_elevation"] = 0
    panel.from_project_dict(json.loads(json.dumps(saved)), skip_dirty_signal=True)
    assert panel.node_table.item(1, 41).text() == "闸前 100.000\n闸后 99.900"
    assert panel._read_table_numeric_value(1, 37) == 0.1
    assert panel._read_table_numeric_value(2, 41) == pytest.approx(99.7)
    assert panel.node_table.item(1, 42).text() == "98.700"
    assert panel.node_table.item(1, 43).text() == "101.200"
    assert panel._build_nodes_from_table()[1].terminal_gate_backfill_details["donor_name"] == "上游明渠"


def test_first_gate_unknown_before_and_explicit_zero_after_remain_distinct(panel):
    gate = ChannelNode(name="起算节制闸", structure_type=StructureType.CHECK_GATE, is_diversion_gate=True)
    gate.water_level = 0
    gate.gate_downstream_water_level = 0
    panel.calculated_nodes = [gate]
    panel._update_table_from_nodes_full([gate])
    assert panel.node_table.item(0, 41).text() == "闸前 —\n闸后 0.000"
    decoded = panel._build_nodes_from_table()[0]
    assert decoded.gate_upstream_water_level is None
    assert decoded.gate_downstream_water_level == 0


def test_gate_detail_dialog_explains_both_sides(panel, monkeypatch):
    from app_渠系计算前端.water_profile import formula_dialog
    captured = []
    monkeypatch.setattr(formula_dialog, "FormulaDialog", lambda _parent, _title, sections: captured.extend(sections))
    panel._on_node_cell_double_clicked(1, 41)
    text = json.dumps(captured, ensure_ascii=False)
    assert "闸前水位" in text and "闸后水位" in text
    assert "99.9000" in text and "99.8000" in text
    assert "来水区间损失" in text and "100.000000 - 0.100000" in text


def test_missing_gate_reference_and_nonfinite_elevations_remain_dash(panel):
    for node in panel.calculated_nodes:
        node.structure_height = 0
    gate = panel.calculated_nodes[1]
    gate.bottom_elevation = float("inf")
    gate.top_elevation = float("nan")
    panel._update_table_from_nodes_full(panel.calculated_nodes)
    assert panel.node_table.item(1, 42).text() == "-"
    assert panel.node_table.item(1, 43).text() == "-"
    panel._refresh_gate_elevation_cells()
    assert panel.node_table.item(1, 42).text() == "-"
    assert panel.node_table.item(1, 43).text() == "-"
    decoded = panel._build_nodes_from_table()[1]
    assert decoded.bottom_elevation == 0
    assert decoded.top_elevation == 0
    assert decoded.terminal_gate_backfill_details["status"] == "failed"


def test_gate_elevation_dialogs_show_reference_and_values(panel, monkeypatch):
    from app_渠系计算前端.water_profile import formula_dialog
    captured = []
    monkeypatch.setattr(formula_dialog, "FormulaDialog", lambda _parent, title, sections: captured.append((title, sections)))
    panel._on_node_cell_double_clicked(1, 42)
    panel._on_node_cell_double_clicked(1, 43)
    text = json.dumps(captured, ensure_ascii=False)
    assert "上游明渠" in text and "同流量段" in text
    assert "无设计台阶" in text and "参考底坡" in text
    assert "98.7000" in text and "101.2000" in text
    assert "末尾闸" not in text


def test_known_zero_gate_bottom_survives_edit_undo_and_project_reopen(panel):
    panel._settings.start_water_level = 1.3
    panel.start_wl_edit.setText("1.3")
    WaterProfileCalculator(panel._settings).calculate_all(panel.calculated_nodes)
    panel._update_table_from_nodes_full(panel.calculated_nodes)
    assert panel.node_table.item(1, 42).text() == "0.000"
    assert panel.node_table.item(1, 43).text() == "2.500"
    panel.node_table.setCurrentCell(1, 37)
    panel.node_table.item(1, 37).setText("0.1250")
    assert panel.node_table.item(1, 42).text() == "0.000"
    assert panel.node_table.item(1, 43).text() == "2.500"
    panel._undo_loss_edit()
    assert panel.calculated_nodes[1].bottom_elevation == 0
    assert panel.node_table.item(1, 42).text() == "0.000"
    panel._redo_loss_edit()
    assert panel.calculated_nodes[1].bottom_elevation == 0
    assert panel.node_table.item(1, 43).text() == "2.500"
    panel._undo_loss_edit()
    saved = json.loads(json.dumps(panel.to_project_dict()))
    panel.from_project_dict(saved, skip_dirty_signal=True)
    assert panel.calculated_nodes[1].bottom_elevation == 0
    assert panel.node_table.item(1, 42).text() == "0.000"
    assert panel.node_table.item(1, 43).text() == "2.500"
    assert panel._build_nodes_from_table()[1].terminal_gate_backfill_details["bottom"]["success"]


def test_excel_gate_water_cell_retains_two_visible_lines(panel, monkeypatch):
    import openpyxl
    captured = []
    monkeypatch.setattr(panel_module.QFileDialog, "getSaveFileName", lambda *args, **kwargs: ("gate.xlsx", ""))
    monkeypatch.setattr(openpyxl.Workbook, "save", lambda wb, _path: captured.append(wb))
    monkeypatch.setattr(panel_module, "ask_open_file", lambda *args, **kwargs: None)
    panel._export_excel()
    sheet = captured[0].active
    cell = sheet.cell(row=5, column=42)
    assert cell.value == "闸前 99.900\n闸后 99.800"
    assert cell.alignment.wrap_text
    assert sheet.row_dimensions[5].height >= 30
    assert sheet.cell(row=5, column=43).value == "98.700"
    assert sheet.cell(row=5, column=44).value == "101.200"


def test_computed_zero_gate_upstream_survives_interval_loss_and_project_reopen(panel):
    # 起点 0.1 m 抵消上游 100 m、底坡 0.001 的区间摩阻，构造真实计算得到的零闸前。
    panel._settings.start_water_level = 0.1
    panel.start_wl_edit.setText("0.1")
    WaterProfileCalculator(panel._settings).calculate_all(panel.calculated_nodes)
    panel._update_table_from_nodes_full(panel.calculated_nodes)
    assert panel._read_table_numeric_value(0, 41) == 0.1
    assert panel._read_table_numeric_value(1, 35) == pytest.approx(0.1)
    assert panel.node_table.item(1, 41).text() == "闸前 0.000\n闸后 -0.100"
    saved = json.loads(json.dumps(panel.to_project_dict()))
    panel.from_project_dict(saved, skip_dirty_signal=True)
    assert panel.node_table.item(1, 41).text() == "闸前 0.000\n闸后 -0.100"
    assert panel._build_nodes_from_table()[1].gate_upstream_water_level == 0


def test_legacy_zero_gate_water_placeholder_is_not_presented_as_known_level(panel):
    first = panel.node_table.item(1, 0)
    payload = first.data(Qt.UserRole)
    payload.pop('_gate_upstream_water_level')
    payload.pop('_gate_downstream_water_level')
    payload['_numeric_display_values']['41'] = {'text': '0.000', 'value': 0}
    first.setData(Qt.UserRole, payload)
    panel.node_table.item(1, 41).setText('0.000')
    panel._refresh_gate_water_level_cells()
    assert panel.node_table.item(1, 41).text() == "闸前 —\n闸后 —"
    assert panel.calculated_nodes[1].gate_upstream_water_level is None
    assert panel.calculated_nodes[1].gate_downstream_water_level is None
